"""El reporte de la previsualización y las plantillas que VisaNow debe llenar.

Dos cosas distintas en un mismo archivo:

- **Lo que la migración haría y lo que no puede hacer sola.** Cada excepción con
  el archivo, la hoja y la fila exactos, para poder abrir el Excel ahí mismo.
- **Las plantillas.** Solo de lo que de verdad hace falta y no está en ningún
  archivo. Una plantilla de 852 filas para llenar a mano no la llena nadie: por
  eso van prellenadas y recortadas a lo que bloquea la carga.

El archivo que sale tiene nombres, teléfonos y montos de clientes reales: vive
en `05_Migracion/`, fuera del repositorio, y no se comparte fuera del proyecto.
"""
from __future__ import annotations

import datetime as dt
import pathlib
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.migracion.motor import Plan

DESTINO = pathlib.Path(__file__).resolve().parents[4] / '05_Migracion'

TITULO = Font(bold=True, color='FFFFFF', size=11)
FONDO = PatternFill('solid', fgColor='0F766E')
POR_LLENAR = PatternFill('solid', fgColor='FEF9C3')
ALERTA = Font(bold=True, color='B91C1C')


def _hoja(wb: Workbook, nombre: str, encabezados: list[str], anchos: list[int]):
    h = wb.create_sheet(nombre)
    for c, (titulo, ancho) in enumerate(zip(encabezados, anchos), start=1):
        celda = h.cell(1, c, titulo)
        celda.font, celda.fill = TITULO, FONDO
        celda.alignment = Alignment(vertical='center', wrap_text=True)
        h.column_dimensions[get_column_letter(c)].width = ancho
    h.freeze_panes = 'A2'
    h.auto_filter.ref = f'A1:{get_column_letter(len(encabezados))}1'
    return h


def escribir(plan: Plan, *, corrida_id: int, hoy: dt.date | None = None) -> pathlib.Path:
    hoy = hoy or dt.date.today()
    wb = Workbook()
    wb.remove(wb.active)

    _resumen(wb, plan, corrida_id, hoy)
    _excepciones(wb, plan)
    _plantilla_usuarios(wb)
    _plantilla_ventas_sin_cobro(wb, plan)
    _plantilla_sin_precio(wb, plan)

    DESTINO.mkdir(exist_ok=True)
    ruta = DESTINO / f'Previsualizacion_migracion_{hoy:%Y-%m-%d}.xlsx'
    wb.save(ruta)
    return ruta


# ------------------------------------------------------------------- resumen

def _resumen(wb: Workbook, plan: Plan, corrida_id: int, hoy: dt.date) -> None:
    h = wb.create_sheet('RESUMEN')
    h.column_dimensions['A'].width = 46
    h.column_dimensions['B'].width = 20
    h.column_dimensions['C'].width = 62

    def linea(fila: int, etiqueta, valor=None, nota=None, negrita=False):
        h.cell(fila, 1, etiqueta).font = Font(bold=negrita or valor is None)
        if valor is not None:
            c = h.cell(fila, 2, valor)
            c.font = Font(bold=negrita)
            if isinstance(valor, (int, float)) and abs(valor) > 999:
                c.number_format = '#,##0'
        if nota:
            c = h.cell(fila, 3, nota)
            c.alignment = Alignment(wrap_text=True, vertical='top')
        return fila + 1

    r = 1
    t = h.cell(r, 1, 'PREVISUALIZACIÓN DE LA MIGRACIÓN')
    t.font = Font(bold=True, size=14, color='0F766E')
    r += 1
    r = linea(r, f'Corrida #{corrida_id} · {hoy:%d/%m/%Y}', None,
              'No se escribió ningún dato: esto es lo que la migración haría.')
    r += 1

    r = linea(r, 'LO QUE ENTRARÍA', negrita=True)
    r = linea(r, 'Clientes', plan.clientes,
              'Personas distintas reconocidas en las 11 hojas. El mismo cliente aparece escrito '
              'de varias formas; aquí ya están unificados.')
    r = linea(r, 'Personas que viajan', plan.solicitantes)
    r = linea(r, 'Trámites', plan.casos)
    r = linea(r, 'Citas (CAS y entrevista)', plan.citas)
    r = linea(r, 'Ventas', plan.negocios)
    r = linea(r, 'Pagos', plan.pagos,
              'Una fila por abono. Hoy son columnas «Abono 1/2/3» y ya se desbordaron.')
    r += 1

    r = linea(r, 'EL DINERO', negrita=True)
    r = linea(r, 'Valor vendido', round(plan.valor_vendido))
    r = linea(r, 'Valor cobrado', round(plan.valor_cobrado))
    r = linea(r, 'CARTERA (lo que quedaría por cobrar)', round(plan.cartera), negrita=True,
              nota='Este número es la prueba de que la migración quedó bien: '
                   'VisaNow lo puede comparar contra su propia cuenta.')
    r += 1

    r = linea(r, 'DE DÓNDE SALE EL DINERO', negrita=True)
    for c, titulo in enumerate(('Hoja', 'Vendido', 'Cobrado', 'Cartera', 'Ventas'), start=1):
        celda = h.cell(r, c, titulo)
        celda.font, celda.fill = TITULO, FONDO
    r += 1
    for nombre, (vendido, cobrado, cuantas) in sorted(
            plan.dinero_por_hoja.items(), key=lambda x: -x[1][0]):
        h.cell(r, 1, nombre.split('::')[1])
        for c, v in enumerate((vendido, cobrado, vendido - cobrado, cuantas), start=2):
            celda = h.cell(r, c, round(v) if c < 5 else v)
            celda.number_format = '#,##0'
        r += 1
    r += 1

    r = linea(r, 'LO QUE NO SE PUEDE MIGRAR SOLO', negrita=True)
    r = linea(r, 'Excepciones', len(plan.excepciones),
              'Cada una con su archivo, hoja y fila en la pestaña EXCEPCIONES.')
    for motivo, cuantas in plan.motivos.most_common(15):
        r = linea(r, f'   {motivo}', cuantas)
    r += 1

    r = linea(r, 'HOJAS QUE NO SE MIGRAN, Y POR QUÉ', negrita=True)
    from app.migracion.fuentes import EXCLUIDAS
    for (archivo, hoja), razon in EXCLUIDAS.items():
        h.cell(r, 1, f'{archivo.split(".")[0][:20]} · {hoja}')
        c = h.cell(r, 3, razon)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        h.row_dimensions[r].height = 30
        r += 1


def _excepciones(wb: Workbook, plan: Plan) -> None:
    h = _hoja(wb, 'EXCEPCIONES',
              ['Archivo', 'Hoja', 'Fila', 'Qué pasa', 'Qué hay que hacer', 'Resuelto (sí/no)'],
              [24, 20, 8, 62, 62, 16])
    for r, e in enumerate(sorted(plan.excepciones, key=lambda x: (x.archivo, x.hoja, x.fila)), start=2):
        for c, v in enumerate((e.archivo, e.hoja, e.fila, e.que, e.que_hacer), start=1):
            celda = h.cell(r, c, v)
            celda.alignment = Alignment(wrap_text=True, vertical='top')
        h.cell(r, 6).fill = POR_LLENAR


# ----------------------------------------------------------------- plantillas

def _plantilla_usuarios(wb: Workbook) -> None:
    h = _hoja(wb, 'PLANTILLA usuarios',
              ['Nombre completo', 'Correo', 'Perfil', '¿Sigue trabajando?', 'Notas'],
              [32, 34, 22, 20, 40])
    h.cell(2, 1, 'Escriba aquí a las personas que van a entrar al sistema.')
    h.cell(2, 1).font = ALERTA
    h.merge_cells('A2:E2')
    # Prellenado con lo que sí se sabe: los nombres que aparecen como vendedores
    # en los archivos. Confirmar es mucho más rápido que escribir desde cero.
    conocidos = [('Laura', '', 'Administradora', '', 'Aparece como vendedora en las hojas de pagos'),
                 ('Angie', '', 'Comercial', '', 'La vendedora con más ventas en los archivos'),
                 ('Yasmín', '', 'Comercial', '', 'Aparece como vendedora y en la hoja de gastos'),
                 ('Miriam', '', 'Apoyo externo', '', 'Aparece en la hoja de mensajería'),
                 ('', '', '', '', 'El primo que lleva el control de la operación')]
    for r, fila in enumerate(conocidos, start=3):
        for c, v in enumerate(fila, start=1):
            h.cell(r, c, v)
            if c in (2, 3, 4):
                h.cell(r, c).fill = POR_LLENAR
    h.cell(9, 1, 'Perfiles disponibles: Administradora · Comercial · Operaciones · Finanzas · '
                 'Apoyo externo · Solo lectura')
    h.merge_cells('A9:E9')


def _plantilla_ventas_sin_cobro(wb: Workbook, plan: Plan) -> None:
    h = _hoja(wb, 'PLANTILLA ventas sin cobro',
              ['Archivo', 'Hoja', 'Fila', 'Qué dice el archivo',
               '¿Se cobró? (sí/no)', '¿Dónde quedó registrado?'],
              [24, 20, 8, 66, 18, 40])
    h.cell(2, 1, 'Estas ventas no tienen ningún pago registrado. Si de verdad no se cobraron, '
                 'entran a la cartera como deuda. Si sí se cobraron, hay que decir dónde está el '
                 'registro para no reclamarle a un cliente que ya pagó.')
    h.cell(2, 1).font = ALERTA
    h.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='top')
    h.merge_cells('A2:F2')
    h.row_dimensions[2].height = 45
    r = 3
    for e in plan.excepciones:
        if 'sin ningún pago registrado' not in e.que and 'no registra abonos' not in e.que:
            continue
        for c, v in enumerate((e.archivo, e.hoja, e.fila, e.que), start=1):
            h.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical='top')
        h.cell(r, 5).fill = POR_LLENAR
        h.cell(r, 6).fill = POR_LLENAR
        r += 1


def _plantilla_sin_precio(wb: Workbook, plan: Plan) -> None:
    h = _hoja(wb, 'PLANTILLA precio pactado',
              ['Archivo', 'Hoja', 'Fila', 'Qué pasa', 'Valor pactado real', '¿Quedó saldo?'],
              [24, 20, 8, 66, 20, 18])
    h.cell(2, 1, 'De estas ventas el archivo solo registra el pago, no cuánto se había pactado. '
                 'La migración asume que el servicio se pagó completo. Corrija las que tenían '
                 'saldo pendiente; las demás se pueden dejar en blanco.')
    h.cell(2, 1).font = ALERTA
    h.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='top')
    h.merge_cells('A2:F2')
    h.row_dimensions[2].height = 45
    r = 3
    for e in plan.excepciones:
        if 'no el valor pactado' not in e.que and 'sin valor pactado' not in e.que:
            continue
        for c, v in enumerate((e.archivo, e.hoja, e.fila, e.que), start=1):
            h.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical='top')
        h.cell(r, 5).fill = POR_LLENAR
        h.cell(r, 6).fill = POR_LLENAR
        r += 1
