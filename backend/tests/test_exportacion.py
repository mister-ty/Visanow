"""Exportación de listas con registro en la auditoría (RF-075, RNF-07).

Estas pruebas NO se corrieron al escribirlas (el entorno no tenía PostgreSQL):
revisar el resultado antes de fusionar.
"""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.esquema import Auditoria
from app.services import usuarios as servicio_usuarios
from tests.conftest import entrar

EXPORTAR = '/api/v1/exportaciones'


@pytest.fixture
def admin(cliente, usuario):
    return entrar(cliente, usuario('administradora'))


def registros(db, entidad):
    return list(db.scalars(select(Auditoria).where(Auditoria.entidad == entidad)))


def test_exporta_clientes_y_deja_registro_de_quien_y_cuantos(cliente, admin, db):
    for n in ('Ana Prueba', 'Beto Prueba'):
        r = cliente.post('/api/v1/clientes', headers=admin, json={'nombre': n})
        assert r.status_code == 201, r.text
    r = cliente.get(f'{EXPORTAR}/clientes', headers=admin, params={'formato': 'csv'})
    assert r.status_code == 200, r.text
    assert 'Ana Prueba' in r.content.decode('utf-8-sig')
    (fila,) = registros(db, 'exportacion:clientes')
    assert fila.operacion == 'exportar'
    assert fila.usuario_id is not None
    assert fila.despues['filas'] >= 2 and fila.despues['formato'] == 'csv'


def test_el_registro_guarda_el_conteo_y_nunca_el_contenido(cliente, admin, db):
    cliente.post('/api/v1/clientes', headers=admin, json={'nombre': 'Nombre Secreto'})
    cliente.get(f'{EXPORTAR}/clientes', headers=admin)
    (fila,) = registros(db, 'exportacion:clientes')
    assert 'Nombre Secreto' not in str(fila.despues)


def test_exporta_a_excel(cliente, admin):
    r = cliente.get(f'{EXPORTAR}/cartera', headers=admin, params={'formato': 'xlsx'})
    assert r.status_code == 200, r.text
    assert 'Cartera' in load_workbook(io.BytesIO(r.content)).sheetnames


def test_un_nombre_que_empieza_por_igual_no_se_ejecuta_como_formula(cliente, admin):
    cliente.post('/api/v1/clientes', headers=admin, json={'nombre': '=HYPERLINK("x")'})
    r = cliente.get(f'{EXPORTAR}/clientes', headers=admin, params={'formato': 'csv'})
    assert "'=HYPERLINK" in r.content.decode('utf-8-sig')


def test_operaciones_no_exporta_cartera_ni_comisiones(cliente, usuario, db):
    """Operaciones no ve plata (29/09): ni en pantalla ni en un archivo."""
    cab = entrar(cliente, usuario('operaciones'))
    for lista in ('cartera', 'comisiones'):
        assert cliente.get(f'{EXPORTAR}/{lista}', headers=cab).status_code == 403
    assert registros(db, 'exportacion:cartera') == []


def test_poder_ver_no_es_poder_exportar(cliente, usuario):
    """Comercial ve clientes pero no tiene clientes.exportar."""
    cab = entrar(cliente, usuario('comercial'))
    assert cliente.get('/api/v1/clientes', headers=cab).status_code == 200
    assert cliente.get(f'{EXPORTAR}/clientes', headers=cab).status_code == 403


def test_finanzas_exporta_cartera_y_comisiones_pero_no_clientes(cliente, usuario):
    cab = entrar(cliente, usuario('finanzas'))
    assert cliente.get(f'{EXPORTAR}/cartera', headers=cab).status_code == 200
    assert cliente.get(f'{EXPORTAR}/comisiones', headers=cab).status_code == 200
    assert cliente.get(f'{EXPORTAR}/clientes', headers=cab).status_code == 403


def test_quien_solo_ve_lo_suyo_no_se_lleva_la_lista_de_toda_la_agencia(cliente, usuario, db):
    admin = usuario('administradora')
    otro = usuario('finanzas')
    servicio_usuarios.editar(db, admin.u, otro.u.id, alcance='propios')
    db.commit()
    cab = entrar(cliente, otro)
    assert cliente.get(f'{EXPORTAR}/cartera', headers=cab).status_code == 403


def test_lista_desconocida_se_rechaza(cliente, admin):
    assert cliente.get(f'{EXPORTAR}/usuarios', headers=admin).status_code == 422


def test_exportar_un_tablero_tambien_queda_registrado(cliente, usuario, db):
    cab = entrar(cliente, usuario('finanzas'))
    r = cliente.get('/api/v1/tableros/financiero/exportar', headers=cab, params={'formato': 'csv'})
    assert r.status_code == 200, r.text
    assert len(registros(db, 'exportacion:tablero-financiero')) == 1
