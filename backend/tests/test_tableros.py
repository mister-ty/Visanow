"""Tableros comercial, operativo, financiero y ejecutivo (RF-070 a RF-075).

Estas pruebas NO se corrieron al escribirlas (el entorno no tenía PostgreSQL):
revisar el resultado antes de fusionar.
"""
import io

import pytest
from openpyxl import load_workbook

from app.services import usuarios as servicio_usuarios
from tests.conftest import entrar

TABLEROS = ('comercial', 'operativo', 'financiero', 'ejecutivo')


@pytest.fixture
def finanzas(cliente, usuario):
    return entrar(cliente, usuario('finanzas'))


@pytest.mark.parametrize('nombre', TABLEROS)
def test_cada_tablero_responde_aunque_no_haya_datos(cliente, finanzas, nombre):
    r = cliente.get(f'/api/v1/tableros/{nombre}', headers=finanzas)
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo['tablero'] == nombre and cuerpo['indicadores']


def test_rango_de_fechas_invertido_se_rechaza(cliente, finanzas):
    r = cliente.get('/api/v1/tableros/comercial', headers=finanzas,
                    params={'desde': '2026-10-31', 'hasta': '2026-10-01'})
    assert r.status_code == 422


def test_exporta_a_excel_y_csv_con_los_mismos_filtros(cliente, finanzas):
    r = cliente.get('/api/v1/tableros/financiero/exportar', headers=finanzas,
                    params={'formato': 'xlsx', 'desde': '2026-01-01'})
    assert r.status_code == 200, r.text
    libro = load_workbook(io.BytesIO(r.content))
    assert 'Indicadores' in libro.sheetnames
    r = cliente.get('/api/v1/tableros/comercial/exportar', headers=finanzas,
                    params={'formato': 'csv'})
    assert r.status_code == 200 and r.content.startswith('﻿'.encode())


def test_quien_solo_ve_lo_suyo_no_abre_los_tableros_de_toda_la_agencia(
        cliente, usuario, db):
    """RNF-03: el alcance es un control de acceso, no un filtro de la pantalla."""
    admin = usuario('administradora')
    otro = usuario('operaciones')
    servicio_usuarios.editar(db, admin.u, otro.u.id, alcance='asignados')
    db.commit()
    cab = entrar(cliente, otro)
    assert cliente.get('/api/v1/tableros/financiero', headers=cab).status_code == 403
    assert cliente.get('/api/v1/tableros/ejecutivo', headers=cab).status_code == 403
    assert cliente.get('/api/v1/tableros/operativo', headers=cab).status_code == 200


# --------------------------------------------- la contabilidad y quien no la ve

def test_operaciones_no_ve_los_tableros_de_plata(cliente, usuario):
    """Respuesta de la administradora del 29/09: a operaciones no hay que
    ocultarle nada «salvo lo de la contabilidad».

    Esto se colaba por una razon sutil: el tablero se restringia por el ALCANCE
    del usuario, y operaciones se crea con alcance «todos» porque tiene que ver
    todos los tramites. Con eso entraba al financiero y veia vendido, recaudado,
    cartera y vencido. Ver plata es un permiso, no un alcance.
    """
    cab = entrar(cliente, usuario('operaciones'))

    for tablero in ('financiero', 'ejecutivo'):
        r = cliente.get(f'/api/v1/tableros/{tablero}', headers=cab)
        assert r.status_code == 403, f'{tablero}: {r.status_code} {r.text[:200]}'
        assert r.json()['codigo'] == 'sin_permiso'

    # Los suyos si los ve: no se le cierra lo que si es su trabajo.
    for tablero in ('comercial', 'operativo'):
        r = cliente.get(f'/api/v1/tableros/{tablero}', headers=cab)
        assert r.status_code == 200, f'{tablero}: {r.status_code} {r.text[:200]}'


def test_finanzas_si_ve_los_tableros_de_plata(cliente, usuario):
    """El candado no puede dejar por fuera a quien si lleva la contabilidad."""
    cab = entrar(cliente, usuario('finanzas'))
    for tablero in ('financiero', 'ejecutivo'):
        r = cliente.get(f'/api/v1/tableros/{tablero}', headers=cab)
        assert r.status_code == 200, f'{tablero}: {r.status_code} {r.text[:200]}'
    claves = {i['clave'] for i in cliente.get('/api/v1/tableros/financiero',
                                              headers=cab).json()['indicadores']}
    assert {'vendido', 'recaudado', 'cartera'} <= claves, claves
